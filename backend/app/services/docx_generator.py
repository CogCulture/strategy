from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from typing import Optional
import io

def add_section_heading(doc: Document, text: str, level: int = 1):
    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

def add_body_text(doc: Document, text: str):
    para = doc.add_paragraph(text)
    para.paragraph_format.space_after = Pt(8)

def generate_output_docx(brand_name: str, outputs: dict, modifications: Optional[list[dict]] = None) -> bytes:
    doc = Document()
    
    # Apply any chat-driven modifications before rendering
    if modifications:
        from app.services.modification_parser import apply_modifications
        outputs = apply_modifications(outputs, modifications)
    
    # Title
    title = doc.add_heading(f"{brand_name} — Brand Research Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")
    
    # Section order
    section_map = {
        "brand_strategy": "Brand Strategy",
        "competition_scan": "Competition Scan",
        "brand_audit": "Brand Audit",
        "positioning": "Positioning",
        "social_media": "Social Media Analysis",
        "seo_audit": "SEO Audit",
        "launch_plan": "Launch Plan",
        "content_buckets": "Content Strategy — Buckets",
        "content_tone": "Content Strategy — Tone & Guidelines",
        "content_messaging": "Content Strategy — Key Messaging",
    }
    
    for key, section_title in section_map.items():
        if key in outputs and outputs[key]:
            add_section_heading(doc, section_title, level=1)
            doc.add_paragraph("")
            
            content = outputs[key]
            if isinstance(content, (dict, list)):
                import json
                content = json.dumps(content, indent=2)
            
            # Split by markdown headers and render
            for line in content.split("\n"):
                if line.startswith("## "):
                    add_section_heading(doc, line[3:], level=2)
                elif line.startswith("### "):
                    add_section_heading(doc, line[4:], level=3)
                elif line.strip():
                    add_body_text(doc, line)
            
            doc.add_page_break()
    
    # Render any dynamically added sections not in the standard map
    for key, content in outputs.items():
        if key not in section_map and content:
            section_title = key.replace("_", " ").title()
            add_section_heading(doc, section_title, level=1)
            doc.add_paragraph("")
            
            if isinstance(content, (dict, list)):
                import json
                content = json.dumps(content, indent=2)
            
            for line in content.split("\n"):
                if line.startswith("## "):
                    add_section_heading(doc, line[3:], level=2)
                elif line.startswith("### "):
                    add_section_heading(doc, line[4:], level=3)
                elif line.strip():
                    add_body_text(doc, line)
            
            doc.add_page_break()
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()
